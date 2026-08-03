"""Index Control SaaS API tests."""

from __future__ import annotations

import uuid
from io import BytesIO
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

pytest.importorskip("fastapi")
from fastapi.testclient import TestClient

from app.core.security import create_access_token, hash_password
from app.db import Base, SessionLocal, engine
from app.models.user import User
from app.services.index_control_history_store import create_check_session
from main import app

client = TestClient(app)


@pytest.fixture()
def auth_headers():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(
            email=f"saas-{uuid.uuid4().hex[:10]}@example.com",
            password_hash=hash_password("test-pass"),
            role="user",
            status="active",
        )
        db.add(user)
        db.commit()
        db.refresh(user)
        token = create_access_token(user_id=user.id)
        yield {"Authorization": f"Bearer {token}"}
    finally:
        db.close()


def test_saas_overview_requires_auth():
    assert client.get("/api/index-control/saas-overview").status_code == 401


def test_saas_overview_ok(auth_headers):
    r = client.get("/api/index-control/saas-overview", headers=auth_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert "quota" in data
    assert "recent_submits" in data


@patch("app.services.index_control_saas_service.sinbyte_submit_urls")
@patch("app.services.index_control_saas_service.resolve_api_key", return_value="abcd1234secret")
def test_submit_indexing_with_project_name(mock_resolve, mock_submit, auth_headers, monkeypatch):
    monkeypatch.setattr(
        "app.services.url_safety.socket.getaddrinfo",
        lambda host, port: [(2, 1, 6, "", ("93.184.216.34", 0))],
    )
    mock_submit.return_value = {"ok": True, "http_status": 200, "message": "Sinbyte OK"}

    r = client.post(
        "/api/index-control/submit-indexing",
        headers=auth_headers,
        json={
            "urls": ["https://example.com/new-page"],
            "project_name": "Link guest post tháng 6",
        },
    )
    assert r.status_code == 200
    data = r.json()
    assert data["summary"]["sent"] >= 1
    assert data.get("project_id")


@patch("app.services.index_control_saas_service.sinbyte_submit_urls")
@patch("app.services.index_control_saas_service.resolve_api_key", return_value="abcd1234secret")
def test_submit_indexing(mock_resolve, mock_submit, auth_headers, monkeypatch):
    monkeypatch.setattr(
        "app.services.url_safety.socket.getaddrinfo",
        lambda host, port: [(2, 1, 6, "", ("93.184.216.34", 0))],
    )
    mock_submit.return_value = {"ok": True, "http_status": 200, "message": "Sinbyte OK"}

    r = client.post(
        "/api/index-control/submit-indexing",
        headers=auth_headers,
        json={"urls": ["https://example.com/new-page"]},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["summary"]["sent"] >= 1
    assert data["message"] == "Submit thành công."
    assert "sinbyte" not in r.text.lower()


@patch("app.services.index_control_saas_service.resolve_api_key", return_value=None)
def test_submit_indexing_service_unavailable(mock_resolve, auth_headers):
    r = client.post(
        "/api/index-control/submit-indexing",
        headers=auth_headers,
        json={"urls": ["https://example.com/page"]},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is False
    assert data["error"] == "service_unavailable"
    assert data["message"] == "Hệ thống submit hiện chưa sẵn sàng, vui lòng thử lại sau."
    assert "sinbyte" not in r.text.lower()


@patch("app.services.index_control_saas_service._sleep_between_urls")
@patch("app.services.index_control_saas_service._run_real_index_check")
def test_check_free(mock_check, mock_sleep, auth_headers, monkeypatch):
    mock_check.return_value = {
        "url": "https://example.com/page",
        "query": "site:https://example.com/page",
        "index_status": "INDEX",
        "reason": "Đã tìm thấy URL trên Google.",
        "http_status": "200",
        "indexability_status": "INDEXABLE",
        "checked_at": "2026-06-18T00:00:00+00:00",
        "retry_count": 0,
        "source": "google_site_search",
    }

    r = client.post(
        "/api/index-control/check-free",
        headers=auth_headers,
        json={"urls": ["https://example.com/page"]},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert data["items"][0]["status"] == "indexed"
    assert data["items"][0]["status_label"] == "Index"
    assert data["items"][0]["query"] == "site:https://example.com/page"
    assert "Google Search Console" in data["disclaimer"]
    assert "sinbyte" not in r.text.lower()


@patch("app.services.index_control_saas_service._sleep_between_urls")
@patch("app.services.index_control_saas_service._run_real_index_check")
def test_check_free_stream_index(mock_check, mock_sleep, auth_headers):
    mock_check.return_value = {
        "url": "https://example.com/page",
        "query": "site:https://example.com/page",
        "index_status": "INDEX",
        "reason": "Đã tìm thấy URL trên Google.",
        "http_status": "200",
        "indexability_status": "INDEXABLE",
        "checked_at": "2026-06-18T00:00:00+00:00",
        "retry_count": 0,
        "source": "google_site_search",
    }

    with client.stream(
        "POST",
        "/api/index-control/check-free/stream",
        headers=auth_headers,
        json={"urls": ["https://example.com/page"], "phase": "index"},
    ) as r:
        assert r.status_code == 200
        body = "".join(r.iter_text())

    assert "event:init" in body
    assert "event:row" in body
    assert "event:progress" in body
    assert "event:done" in body
    assert "\"INDEX\"" in body
    assert "sinbyte" not in body.lower()


@patch("app.services.index_control_saas_service._sleep_between_urls")
@patch("app.services.index_control_saas_service._run_real_index_check")
def test_check_free_stream_with_project_name(mock_check, mock_sleep, auth_headers):
    mock_check.return_value = {
        "url": "https://example.com/page",
        "query": "site:https://example.com/page",
        "index_status": "INDEX",
        "reason": "OK",
        "http_status": "200",
        "indexability_status": "INDEXABLE",
        "checked_at": "2026-06-18T00:00:00+00:00",
        "retry_count": 0,
        "source": "google_site_search",
    }

    with client.stream(
        "POST",
        "/api/index-control/check-free/stream",
        headers=auth_headers,
        json={
            "urls": ["https://example.com/page"],
            "phase": "index",
            "project_name": "Check guest post tháng 6",
        },
    ) as r:
        assert r.status_code == 200
        body = "".join(r.iter_text())

    assert "event:done" in body
    assert "check_id" in body


@patch("app.services.index_control_saas_service.check_url_response_details")
def test_check_free_stream_response(mock_resp, auth_headers):
    mock_resp.return_value = {
        "url": "https://example.com/page",
        "http_status": "200",
        "indexability": "Có thể index",
        "page_title": "Example Page",
        "message": "",
        "checked_at": "2026-06-18T00:00:00+00:00",
    }

    with client.stream(
        "POST",
        "/api/index-control/check-free/stream",
        headers=auth_headers,
        json={"urls": ["https://example.com/page"], "phase": "response"},
    ) as r:
        assert r.status_code == 200
        body = "".join(r.iter_text())

    assert "event:row" in body
    assert "Example Page" in body
    assert "Có thể index" in body


def test_submit_projects_list(auth_headers):
    r = client.get("/api/index-control/submit-indexing/projects", headers=auth_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert "projects" in data


def test_submit_projects_delete_all(auth_headers):
    r = client.delete("/api/index-control/submit-indexing/projects", headers=auth_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert "deleted" in data


def test_check_free_sessions(auth_headers):
    r = client.get("/api/index-control/check-free/sessions", headers=auth_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert "sessions" in data


def test_check_free_sessions_delete_all(auth_headers):
    r = client.delete("/api/index-control/check-free/sessions", headers=auth_headers)
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert "deleted" in data


def test_check_free_session_isolation(auth_headers):
    headers_a = auth_headers
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user_b = User(
            email=f"saas-b-{uuid.uuid4().hex[:10]}@example.com",
            password_hash=hash_password("test-pass"),
            role="user",
            status="active",
        )
        db.add(user_b)
        db.commit()
        db.refresh(user_b)
        token_b = create_access_token(user_id=user_b.id)
        check_id = create_check_session(
            db,
            user_id=user_b.id,
            domain="example.com",
            url_type="check_free",
            project_name="Private batch",
        )
    finally:
        db.close()

    r = client.get(f"/api/index-control/check-free/sessions/{check_id}", headers=headers_a)
    assert r.status_code == 404
    rb = client.get(
        f"/api/index-control/check-free/sessions/{check_id}",
        headers={"Authorization": f"Bearer {token_b}"},
    )
    assert rb.status_code == 200


def test_check_free_export_rows(auth_headers):
    r = client.post(
        "/api/index-control/check-free/export",
        headers=auth_headers,
        json={
            "rows": [
                {
                    "url": "https://example.com/page",
                    "status": "indexed",
                    "status_label": "Index",
                    "http_status": "200",
                    "indexability": "Có thể index",
                    "page_title": "Demo",
                    "message": "OK",
                }
            ]
        },
    )
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    assert len(r.content) > 100
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    headers = [ws.cell(row=1, column=i).value for i in range(1, 12)]
    assert headers == [
        "STT",
        "URL",
        "Query site",
        "Index status",
        "HTTP status",
        "Khả năng index",
        "Matched URL",
        "Title",
        "Reason",
        "Checked at",
        "Retry count",
    ]


@patch("app.services.index_control_saas_service._sleep_between_urls")
@patch("app.services.index_control_saas_service._run_real_index_check")
def test_check_free_timeout_does_not_crash_batch(mock_check, mock_sleep, auth_headers):
    mock_check.side_effect = [
        {
            "url": "https://example.com/a",
            "query": "site:https://example.com/a",
            "index_status": "PENDING_RETRY",
            "reason": "Google search timeout — thử lại sau.",
            "http_status": "200",
            "indexability_status": "INDEXABLE",
            "checked_at": "2026-06-18T00:00:00+00:00",
            "retry_count": 1,
            "source": "google_site_search",
        },
        {
            "url": "https://example.com/b",
            "query": "site:https://example.com/b",
            "index_status": "NO_INDEX",
            "reason": "Không thấy URL khớp trong kết quả Google.",
            "http_status": "200",
            "indexability_status": "INDEXABLE",
            "checked_at": "2026-06-18T00:00:01+00:00",
            "retry_count": 0,
            "source": "google_site_search",
        },
    ]

    r = client.post(
        "/api/index-control/check-free",
        headers=auth_headers,
        json={"urls": ["https://example.com/a", "https://example.com/b"]},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert len(data["items"]) == 2
    assert data["items"][0]["index_status"] == "PENDING_RETRY"
    assert data["items"][1]["index_status"] == "NO_INDEX"


@patch("app.services.submit_project_store.check_urls_index_status")
@patch("app.services.index_control_saas_service.sinbyte_submit_urls")
@patch("app.services.index_control_saas_service.resolve_api_key", return_value="abcd1234secret")
def test_submit_project_check_index(mock_resolve, mock_submit, mock_check, auth_headers, monkeypatch):
    monkeypatch.setattr(
        "app.services.url_safety.socket.getaddrinfo",
        lambda host, port: [(2, 1, 6, "", ("93.184.216.34", 0))],
    )
    mock_submit.return_value = {"ok": True, "http_status": 200, "message": "OK"}
    mock_check.return_value = [
        {
            "url": "https://example.com/new-page",
            "index_status": "indexed",
            "message": "Indexed",
            "source": "technical",
            "checked_at": "2026-06-18T00:00:00+00:00",
        }
    ]

    submit = client.post(
        "/api/index-control/submit-indexing",
        headers=auth_headers,
        json={
            "urls": ["https://example.com/new-page"],
            "project_name": "Entity test project",
        },
    )
    project_id = submit.json().get("project_id")
    assert project_id

    r = client.post(
        f"/api/index-control/submit-indexing/projects/{project_id}/check-index",
        headers=auth_headers,
    )
    assert r.status_code == 200
    data = r.json()
    assert data["ok"] is True
    assert data["items"]
    assert data["items"][0]["indexed"] == "Index"
