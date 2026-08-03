"""
Parse clustered keyword exports into Content AI bulk preview (1 article / cluster).

Supports:
- Flat export: Từ khóa chính | Từ khóa phụ | Phân Loại | Volumg seach
- English: main_keyword | sub_keyword | intent | search_volume
- Row format with cluster_id + primary_keyword / is_primary
- Clusters sheet: cluster_id | main_keyword | intent | keywords
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from io import BytesIO, StringIO
from typing import Any

_MAX_CLUSTERS = 200
_MAX_MEMBERS = 5000

_CLUSTER_ID_ALIASES = ("cluster_id", "cluster", "nhóm", "nhom", "group_id", "group", "cụm", "cum")
_PRIMARY_ALIASES = (
    "primary_keyword",
    "main_keyword",
    "từ khóa chính",
    "tu khoa chinh",
    "keyword chính",
    "keyword chinh",
)
_SUB_ALIASES = (
    "sub_keyword",
    "secondary_keyword",
    "từ khóa phụ",
    "tu khoa phu",
    "keyword phụ",
    "keyword phu",
)
_KEYWORD_ALIASES = ("keyword", "từ khóa", "tu khoa", "kw", "key word")
_VOLUME_ALIASES = (
    "search_volume",
    "volume",
    "volumg seach",
    "vol",
    "search vol",
    "search volume",
)
_INTENT_ALIASES = ("intent", "phân loại", "phan loai", "classification", "search_intent")
_IS_PRIMARY_ALIASES = ("is_primary", "primary", "là chính", "la chinh", "is primary")


def _norm_header(value: Any) -> str:
    s = str(value or "").strip().lower()
    return re.sub(r"\s+", " ", s)


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _match_alias(header: str, aliases: tuple[str, ...]) -> bool:
    h = _norm_header(header)
    if not h:
        return False
    for a in aliases:
        if h == a:
            return True
    # Chỉ khi header CHỨA alias (không ngược lại) — tránh «keyword» khớp «keyword chính»
    for a in aliases:
        if len(a) >= 8 and a in h:
            return True
    return False


def _map_cluster_headers(headers: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(headers):
        h = _cell_str(cell)
        if not h:
            continue
        if "cluster_id" not in mapping and _match_alias(h, _CLUSTER_ID_ALIASES):
            mapping["cluster_id"] = idx
        elif "primary_keyword" not in mapping and _match_alias(h, _PRIMARY_ALIASES):
            mapping["primary_keyword"] = idx
        elif "sub_keyword" not in mapping and _match_alias(h, _SUB_ALIASES):
            mapping["sub_keyword"] = idx
        elif "keyword" not in mapping and _match_alias(h, _KEYWORD_ALIASES):
            mapping["keyword"] = idx
        elif "search_volume" not in mapping and _match_alias(h, _VOLUME_ALIASES):
            mapping["search_volume"] = idx
        elif "intent" not in mapping and _match_alias(h, _INTENT_ALIASES):
            mapping["intent"] = idx
        elif "is_primary" not in mapping and _match_alias(h, _IS_PRIMARY_ALIASES):
            mapping["is_primary"] = idx
    # Nếu vừa có primary vừa có keyword trùng cột — giữ primary, bỏ keyword
    if (
        "primary_keyword" in mapping
        and "keyword" in mapping
        and mapping["primary_keyword"] == mapping["keyword"]
    ):
        mapping.pop("keyword", None)
    return mapping


def _truthy_primary(raw: str) -> bool:
    s = _norm_header(raw)
    return s in {"1", "true", "yes", "y", "primary", "chinh", "chính", "main", "x"}


def _parse_volume(raw: Any) -> int:
    try:
        return max(0, int(float(str(raw or "0").replace(",", "").strip() or "0")))
    except (TypeError, ValueError):
        return 0


def _looks_like_clustered_headers(headers: list[Any]) -> bool:
    m = _map_cluster_headers(headers)
    if "cluster_id" in m:
        return True
    if "primary_keyword" in m and ("sub_keyword" in m or "keyword" in m):
        return True
    return False


def _pick_primary_keyword(members: list[dict[str, Any]], *, cluster_intent: str = "") -> str:
    """Ưu tiên volume + độ ngắn tự nhiên + khớp intent."""
    if not members:
        return ""
    marked = [m for m in members if m.get("is_primary")]
    if len(marked) == 1:
        return str(marked[0].get("keyword") or "").strip()
    if len(marked) > 1:
        # nhiều primary → vẫn chọn theo score trong tập marked
        pool = marked
    else:
        pool = members

    max_vol = max(int(m.get("search_volume") or 0) for m in pool) or 1
    best_kw = ""
    best_score = -1.0
    intent = (cluster_intent or "").strip().lower()

    for m in pool:
        kw = str(m.get("keyword") or "").strip()
        if not kw:
            continue
        vol = int(m.get("search_volume") or 0)
        words = kw.split()
        wc = len(words)
        row_intent = str(m.get("intent") or "").strip().lower()
        intent_fit = 1.0 if (intent and row_intent == intent) else (0.7 if not intent else 0.4)
        if 2 <= wc <= 5:
            len_score = 1.0
        elif wc == 1 or wc <= 7:
            len_score = 0.75
        else:
            len_score = 0.45
        vol_score = (vol / max_vol) if vol > 0 else 0.2
        clarity = min(1.0, len({w.lower() for w in words}) / max(1, wc))
        score = vol_score * 0.4 + len_score * 0.25 + intent_fit * 0.2 + clarity * 0.15
        if score > best_score:
            best_score = score
            best_kw = kw
    return best_kw or str(pool[0].get("keyword") or "").strip()


def _dominant_intent(members: list[dict[str, Any]]) -> str:
    counts: dict[str, int] = defaultdict(int)
    for m in members:
        intent = str(m.get("intent") or "").strip().lower()
        if intent:
            counts[intent] += 1
    if not counts:
        return ""
    return max(counts.items(), key=lambda x: x[1])[0]


def _build_cluster_preview(
    *,
    cluster_id: str,
    members: list[dict[str, Any]],
    forced_primary: str = "",
) -> dict[str, Any]:
    warnings: list[str] = []
    cleaned: list[dict[str, Any]] = []
    seen_kw: set[str] = set()
    for m in members:
        kw = re.sub(r"\s+", " ", str(m.get("keyword") or "").strip())
        if not kw:
            continue
        lk = kw.lower()
        if lk in seen_kw:
            continue
        seen_kw.add(lk)
        cleaned.append(
            {
                "keyword": kw,
                "search_volume": int(m.get("search_volume") or 0),
                "intent": str(m.get("intent") or "").strip(),
                "is_primary": bool(m.get("is_primary")),
            }
        )

    if not cleaned:
        return {}

    cid = str(cluster_id or "").strip()
    if not cid:
        warnings.append("Thiếu cluster ID — đã gán ID tạm.")
        cid = f"auto_{cleaned[0]['keyword'][:40]}"

    marked = [m for m in cleaned if m.get("is_primary")]
    if len(marked) > 1:
        warnings.append(f"Nhiều từ khóa chính ({len(marked)}) — hệ thống chọn 1 theo volume/intent.")

    intents = {str(m.get("intent") or "").strip().lower() for m in cleaned if str(m.get("intent") or "").strip()}
    if len(intents) > 1:
        warnings.append("Nhóm có nhiều search intent khác nhau.")

    cluster_intent = _dominant_intent(cleaned)
    primary = re.sub(r"\s+", " ", str(forced_primary or "").strip())
    member_keys = {m["keyword"].lower() for m in cleaned}
    if primary and primary.lower() not in member_keys:
        warnings.append("Từ khóa chính chỉ định không thuộc nhóm — chọn lại tự động.")
        primary = ""
    if not primary:
        primary = _pick_primary_keyword(cleaned, cluster_intent=cluster_intent)

    secondaries = [m["keyword"] for m in cleaned if m["keyword"].lower() != primary.lower()]
    primary_vol = next(
        (int(m["search_volume"]) for m in cleaned if m["keyword"].lower() == primary.lower()),
        0,
    )

    return {
        "cluster_id": cid,
        "primary_keyword": primary,
        "secondary_keywords": secondaries,
        "secondary_count": len(secondaries),
        "search_volume": primary_vol,
        "intent": cluster_intent,
        "members": cleaned,
        "warnings": warnings,
        "article_count": 1,
    }


def _rows_from_flat_table(table: list[list[Any]]) -> list[dict[str, Any]]:
    if not table:
        return []
    header_map = _map_cluster_headers(table[0])
    if not header_map:
        return []
    start = 1 if any(_norm_header(_cell_str(c)) for c in table[0]) else 0
    out: list[dict[str, Any]] = []
    for row in table[start:]:
        if len(out) >= _MAX_MEMBERS:
            break

        def _get(field: str) -> str:
            idx = header_map.get(field)
            if idx is None or idx >= len(row):
                return ""
            return _cell_str(row[idx])

        primary = _get("primary_keyword")
        sub = _get("sub_keyword")
        keyword = _get("keyword") or sub or primary
        if not keyword:
            continue
        cluster_id = _get("cluster_id")
        if not cluster_id and primary:
            cluster_id = primary
        is_primary = False
        if "is_primary" in header_map:
            is_primary = _truthy_primary(_get("is_primary"))
        elif primary and keyword.lower() == primary.lower():
            is_primary = True
        elif not sub and primary and keyword.lower() == primary.lower():
            is_primary = True
        elif cluster_id and keyword.lower() == cluster_id.lower():
            # Export Keyword|Volume|Group: Group thường = từ khóa chính
            is_primary = True

        out.append(
            {
                "cluster_id": cluster_id or primary or keyword,
                "keyword": keyword,
                "search_volume": _parse_volume(_get("search_volume")),
                "intent": _get("intent"),
                "is_primary": is_primary,
                "declared_primary": primary or (keyword if is_primary and cluster_id and keyword.lower() == cluster_id.lower() else ""),
            }
        )
    return out


def _rows_from_clusters_sheet(table: list[list[Any]]) -> list[dict[str, Any]]:
    """cluster_id | main_keyword | intent | keywords (comma-separated)."""
    if not table:
        return []
    header_map = _map_cluster_headers(table[0])
    # keywords column may be last unlabeled — find "keywords"
    kw_col = None
    for idx, cell in enumerate(table[0]):
        h = _norm_header(_cell_str(cell))
        if h in {"keywords", "members", "danh sách", "danh sach"}:
            kw_col = idx
            break
    if kw_col is None and len(table[0]) >= 4:
        kw_col = 3

    out: list[dict[str, Any]] = []
    for row in table[1:]:
        if len(out) >= _MAX_MEMBERS:
            break

        def _get(field: str, default_idx: int | None = None) -> str:
            idx = header_map.get(field, default_idx)
            if idx is None or idx >= len(row):
                return ""
            return _cell_str(row[idx])

        cid = _get("cluster_id", 0)
        main = _get("primary_keyword", 1)
        intent = _get("intent", 2)
        joined = ""
        if kw_col is not None and kw_col < len(row):
            joined = _cell_str(row[kw_col])
        members = [x.strip() for x in re.split(r"[,;\n]+", joined) if x.strip()]
        if main and main not in members:
            members.insert(0, main)
        if not members and main:
            members = [main]
        for kw in members:
            out.append(
                {
                    "cluster_id": cid or main or kw,
                    "keyword": kw,
                    "search_volume": 0,
                    "intent": intent,
                    "is_primary": kw.lower() == (main or "").lower(),
                    "declared_primary": main,
                }
            )
    return out


def _group_member_rows(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    global_warnings: list[str] = []
    by_cluster: dict[str, list[dict[str, Any]]] = defaultdict(list)
    missing_id = 0
    for r in rows:
        cid = str(r.get("cluster_id") or "").strip()
        if not cid:
            missing_id += 1
            cid = f"_noid_{str(r.get('keyword') or '').lower()}"
        by_cluster[cid].append(r)

    if missing_id:
        global_warnings.append(f"{missing_id} dòng thiếu cluster ID.")

    # duplicate keywords across clusters
    kw_owners: dict[str, list[str]] = defaultdict(list)
    for cid, mems in by_cluster.items():
        seen: set[str] = set()
        for m in mems:
            lk = str(m.get("keyword") or "").strip().lower()
            if not lk or lk in seen:
                continue
            seen.add(lk)
            kw_owners[lk].append(cid)
    dups = [kw for kw, owners in kw_owners.items() if len(set(owners)) > 1]
    if dups:
        global_warnings.append(
            f"{len(dups)} từ khóa xuất hiện ở nhiều nhóm (vd: {', '.join(dups[:3])}{'…' if len(dups) > 3 else ''})."
        )

    clusters: list[dict[str, Any]] = []
    for cid, mems in by_cluster.items():
        if len(clusters) >= _MAX_CLUSTERS:
            global_warnings.append(f"Đã giới hạn {_MAX_CLUSTERS} nhóm.")
            break
        # Prefer declared primary from flat rows
        forced = ""
        for m in mems:
            dp = str(m.get("declared_primary") or "").strip()
            if dp:
                forced = dp
                break
        preview = _build_cluster_preview(cluster_id=cid if not cid.startswith("_noid_") else "", members=mems, forced_primary=forced)
        if preview:
            clusters.append(preview)

    return clusters, global_warnings


def clusters_to_bulk_items(clusters: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Convert preview clusters → bulk import items (1 row = 1 article)."""
    items: list[dict[str, Any]] = []
    seen: set[str] = set()
    for c in clusters or []:
        pk = re.sub(r"\s+", " ", str(c.get("primary_keyword") or "").strip())
        if not pk:
            continue
        lk = pk.lower()
        if lk in seen:
            continue
        seen.add(lk)
        secs = c.get("secondary_keywords") or []
        if isinstance(secs, str):
            sec_list = [x.strip() for x in re.split(r"[,;\n]+", secs) if x.strip()]
        else:
            sec_list = [str(x).strip() for x in secs if str(x).strip()]
        sec_list = [s for s in sec_list if s.lower() != lk]
        intent = str(c.get("intent") or "").strip()
        items.append(
            {
                "keyword": pk,
                "custom_title": "",
                "custom_description": f"Search intent: {intent}" if intent else "",
                "custom_outline": "",
                "search_volume": int(c.get("search_volume") or 0),
                "content_type": "",
                "target_word_count": None,
                "competitor_url": "",
                "secondary_keywords": sec_list,
                "cluster_id": str(c.get("cluster_id") or ""),
                "intent": intent,
            }
        )
    return items


def _read_xlsx_sheets(raw: bytes) -> dict[str, list[list[Any]]]:
    import openpyxl

    wb = openpyxl.load_workbook(filename=BytesIO(raw), data_only=True, read_only=True)
    sheets: dict[str, list[list[Any]]] = {}
    try:
        for ws in wb.worksheets:
            table: list[list[Any]] = []
            max_r = min((ws.max_row or 0) + 1, _MAX_MEMBERS + 5)
            max_c = min((ws.max_column or 0) + 1, 24)
            for r_idx in range(1, max_r):
                row = [ws.cell(row=r_idx, column=c).value for c in range(1, max_c)]
                if any(v is not None and str(v).strip() for v in row):
                    table.append(row)
            if table:
                sheets[str(ws.title or "Sheet")] = table
    finally:
        wb.close()
    return sheets


def parse_clustered_table(table: list[list[Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    if not table:
        return [], ["File trống."]
    headers = table[0]
    header_map = _map_cluster_headers(headers)
    # Clusters sheet style: has keywords column joined
    has_joined = any(
        _norm_header(_cell_str(c)) in {"keywords", "members", "danh sách", "danh sach"} for c in headers
    )
    if has_joined and "primary_keyword" in header_map:
        rows = _rows_from_clusters_sheet(table)
    else:
        rows = _rows_from_flat_table(table)
    if not rows:
        return [], ["Không nhận diện được định dạng file gom nhóm."]
    return _group_member_rows(rows)


def parse_clustered_upload(*, filename: str, raw: bytes) -> dict[str, Any]:
    name = (filename or "").lower().strip()
    fmt = "xlsx"
    clusters: list[dict[str, Any]] = []
    warnings: list[str] = []

    if name.endswith(".xlsx"):
        sheets = _read_xlsx_sheets(raw)
        # Prefer sheet named Clusters / Gom nhóm
        preferred = None
        for title, table in sheets.items():
            t = title.lower()
            if "cluster" in t or "gom" in t or "nhóm" in t or "nhom" in t:
                preferred = table
                break
        if preferred is None:
            # Pick first sheet that looks clustered
            for table in sheets.values():
                if table and _looks_like_clustered_headers(table[0]):
                    preferred = table
                    break
        if preferred is None and sheets:
            preferred = next(iter(sheets.values()))
        clusters, warnings = parse_clustered_table(preferred or [])
    elif name.endswith(".csv") or name.endswith(".txt"):
        fmt = "csv" if name.endswith(".csv") else "text"
        text = raw.decode("utf-8", errors="replace")
        sep = "," if name.endswith(".csv") else ("\t" if "\t" in text.split("\n", 1)[0] else ",")
        reader = csv.reader(StringIO(text), delimiter=sep)
        table = list(reader)
        clusters, warnings = parse_clustered_table(table)
    elif name.endswith(".xls"):
        raise ValueError("File .xls cũ chưa hỗ trợ — vui lòng lưu lại định dạng .xlsx.")
    else:
        raise ValueError("Chỉ hỗ trợ .xlsx, .txt hoặc .csv.")

    if not clusters:
        raise ValueError(
            "Không tìm thấy nhóm từ khóa hợp lệ. Cần cột cluster_id hoặc "
            "«Từ khóa chính» + «Từ khóa phụ» / main_keyword + sub_keyword."
        )

    return {
        "mode": "clustered",
        "preview": True,
        "article_count": len(clusters),
        "clusters": clusters,
        "warnings": warnings,
        "meta": {
            "count": len(clusters),
            "format": fmt,
            "filename": filename,
            "member_count": sum(len(c.get("members") or []) for c in clusters),
        },
    }
