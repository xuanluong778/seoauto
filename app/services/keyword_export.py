"""
Export keyword research + clusters to Excel (.xlsx) or Google Sheets.
"""

from __future__ import annotations

import io
import os
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook

_KW_HEADERS = [
    "keyword",
    "canonical_keyword",
    "intent",
    "intent_source",
    "raw_volume",
    "display_volume",
    "volume_data_source",
    "updated_at",
    "trend_direction",
    "trend_pct_change",
    "trend_data_source",
    "difficulty_kd",
    "kd_label",
    "kd_data_source",
    "word_count",
    "cpc_avg",
    "cpc_min",
    "cpc_max",
    "cpc_data_source",
    "competition",
    "competition_data_source",
    "allintitle_estimate",
    "allintitle_data_source",
    "data_warnings",
]


def _row_export_values(row: dict[str, Any]) -> list[Any]:
    sv = row.get("search_volume") or {}
    raw_vol = row.get("raw_volume")
    if raw_vol is None:
        raw_vol = sv.get("raw_volume")
    display_vol = row.get("display_volume")
    if display_vol is None:
        display_vol = sv.get("display_volume")
    if display_vol is None:
        display_vol = sv.get("avg_monthly")
    if display_vol is None and isinstance(row.get("search_volume"), (int, float)):
        display_vol = row.get("search_volume")
    ds = row.get("data_source") or {}
    trend = row.get("trend") or {}
    warnings = row.get("data_warnings") or []
    return [
        str(row.get("keyword") or ""),
        str(row.get("canonical_keyword") or (row.get("keyword") or "").lower().strip()),
        str(row.get("intent") or "unknown"),
        str(row.get("intent_source") or "N/A"),
        raw_vol if raw_vol is not None else "N/A",
        display_vol if display_vol is not None else "N/A",
        str(ds.get("volume") or sv.get("data_source") or "N/A"),
        str(row.get("updated_at") or sv.get("updated_at") or "N/A"),
        str(trend.get("direction") or "unknown"),
        trend.get("pct_change") if trend.get("pct_change") is not None else "N/A",
        str(trend.get("data_source") or ds.get("trend") or "N/A"),
        row.get("difficulty") if row.get("difficulty") is not None else "N/A",
        str(row.get("kd_label") or "N/A"),
        str(ds.get("kd") or "N/A"),
        row.get("word_count") if row.get("word_count") is not None else "",
        row.get("cpc_avg") if row.get("cpc_avg") is not None else "N/A",
        row.get("cpc_min") if row.get("cpc_min") is not None else "N/A",
        row.get("cpc_max") if row.get("cpc_max") is not None else "N/A",
        str(ds.get("cpc") or "N/A"),
        str(row.get("competition") or "Unknown"),
        str(ds.get("competition") or "N/A"),
        row.get("allintitle_estimate") if row.get("allintitle_estimate") is not None else "N/A",
        str(ds.get("allintitle") or "N/A"),
        ", ".join(str(w) for w in warnings) if warnings else "",
    ]


def export_to_excel(
    *,
    keywords: list[dict[str, Any]],
    clusters: list[dict[str, Any]],
) -> bytes:
    """
    Sheet **Keywords**: full research metrics with data_source + trend.
    Sheet **Clusters**: cluster_id, main_keyword, intent, keywords (comma-separated).
    """
    wb = Workbook()
    ws_kw = wb.active
    ws_kw.title = "Keywords"
    ws_kw.append(_KW_HEADERS)
    for row in keywords or []:
        ws_kw.append(_row_export_values(row))

    ws_cl = wb.create_sheet("Clusters")
    ws_cl.append(["cluster_id", "main_keyword", "intent", "keywords"])
    for c in clusters or []:
        kws = c.get("keywords") or []
        if kws and isinstance(kws[0], dict):
            joined = ", ".join(str(x.get("keyword") or "") for x in kws if x.get("keyword"))
        else:
            joined = ", ".join(str(x) for x in kws if x)
        ws_cl.append(
            [
                str(c.get("cluster_id") or ""),
                str(c.get("main_keyword") or ""),
                str(c.get("intent") or ""),
                joined,
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_cluster_flat_to_excel(*, rows: list[dict[str, Any]]) -> bytes:
    """
    Flat cluster export (1 row = 1 sub-keyword under a main topic).

    Columns (Vietnamese):
    1) Từ khóa chính
    2) Từ khóa phụ
    3) Phân Loại
    4) Volumg seach
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Clusters"
    ws.append(["Từ khóa chính", "Từ khóa phụ", "Phân Loại", "Volumg seach"])
    for r in rows or []:
        ws.append(
            [
                str(r.get("main_keyword") or ""),
                str(r.get("sub_keyword") or ""),
                str(r.get("classification") or ""),
                r.get("search_volume") if r.get("search_volume") is not None else "",
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_to_google_sheets(
    *,
    keywords: list[dict[str, Any]],
    clusters: list[dict[str, Any]],
    title: str | None = None,
) -> tuple[str | None, str | None]:
    """
    Returns ``(spreadsheet_url, error_message)``. Requires service-account JSON path in env.
    """
    path = (os.getenv("GOOGLE_SHEETS_SA_JSON") or os.getenv("GOOGLE_SHEETS_CREDENTIALS_PATH") or "").strip()
    if not path or not os.path.isfile(path):
        return None, "Set GOOGLE_SHEETS_SA_JSON to a service-account JSON file path (Sheets API enabled)."

    try:
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
    except ImportError as e:
        return None, f"google api client missing: {e}"

    scopes = ["https://www.googleapis.com/auth/spreadsheets"]
    creds = service_account.Credentials.from_service_account_file(path, scopes=scopes)
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)

    t = title or f"Keyword export {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')} UTC"
    spreadsheet_body = {
        "properties": {"title": t[:99]},
        "sheets": [
            {"properties": {"title": "Keywords", "gridProperties": {"rowCount": 5000, "columnCount": 28}}},
            {"properties": {"title": "Clusters", "gridProperties": {"rowCount": 2000, "columnCount": 6}}},
        ],
    }
    created = service.spreadsheets().create(body=spreadsheet_body, fields="spreadsheetId,spreadsheetUrl").execute()
    sid = created.get("spreadsheetId")
    surl = created.get("spreadsheetUrl")

    kw_vals: list[list[Any]] = [_KW_HEADERS]
    for row in keywords or []:
        kw_vals.append(_row_export_values(row))

    cl_vals: list[list[Any]] = [["cluster_id", "main_keyword", "intent", "keywords"]]
    for c in clusters or []:
        kws = c.get("keywords") or []
        if kws and isinstance(kws[0], dict):
            joined = ", ".join(str(x.get("keyword") or "") for x in kws if x.get("keyword"))
        else:
            joined = ", ".join(str(x) for x in kws if x)
        cl_vals.append(
            [
                str(c.get("cluster_id") or ""),
                str(c.get("main_keyword") or ""),
                str(c.get("intent") or ""),
                joined,
            ]
        )

    body = {
        "valueInputOption": "RAW",
        "data": [
            {"range": "Keywords!A1", "values": kw_vals},
            {"range": "Clusters!A1", "values": cl_vals},
        ],
    }
    service.spreadsheets().values().batchUpdate(spreadsheetId=sid, body=body).execute()
    return surl, None
