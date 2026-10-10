"""
Keyword Research + Clustering APIs and HTML tool pages.
"""

from __future__ import annotations

import json
import os
import threading
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel, Field
from sqlalchemy import desc, inspect as sqla_inspect, insert as sqla_insert
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session, class_mapper, load_only

from app.db import ensure_db_alive, get_db, release_db_before_long_io
from app.models.keyword_cluster_project import KeywordClusterProject
from app.models.keyword_research_project import KeywordResearchProject
from app.models.user import User
from app.services.api_payload_optimizer import (
    shape_keyword_cluster_history_response,
    shape_keyword_research_response,
)
from app.services.auth import get_billing_user, get_current_user
from app.services.circuit_breaker import CircuitOpenError
from app.services.trial_access import require_active_trial
from app.services.rbac import normalize_role
from app.services.job_store import get_job_for_user
from app.services.user_scope import assert_job_access
from app.services.async_task_registry import assert_task_owner, register_task_owner
from app.services.credits import (
    cost_cluster_async,
    cost_cluster_save,
    cost_cluster_sync,
    cost_import_excel,
    cost_keyword_suggest,
    cost_research_run,
    credits_enforced,
    volume_batch_charge,
)
from app.services.permissions import PERM_WORKSPACE_WRITE, assert_permission
from app.services.tenant_org import personal_organization_id, resolve_organization_id, user_belongs_to_organization
from app.services.cost_control import CostGateRequest, ExternalApiProvider, cost_gate_session
from app.services.plan_quota_registry import feature_key_for
from app.services.plan_quota_service import record_quota_usage
from app.services.keyword_cluster_pipeline import build_keyword_cluster_api_response
from app.services.keyword_export import export_cluster_flat_to_excel, export_to_excel, export_to_google_sheets
from app.services.keyword_research_cost import KeywordResearchCostExceeded
from app.services.keyword_research_pipeline import build_keyword_research_api_response
from app.services.search_volume import enrich_keyword_volumes_cached
from app.services.saas_gate import enforce_saas_access, saas_action_gate, build_job_id
from app.services.saas_trial_service import TRIAL_BLOCKED_FEATURES
import uuid

router = APIRouter(prefix="/keywords", tags=["keywords"])
templates = Jinja2Templates(directory="templates")


def _kw_org_id(db: Session, user: User, organization_id: int | None = None) -> int:
    from app.services.tenant_org import OrganizationForbiddenError

    try:
        return resolve_organization_id(int(user.id), organization_id, db=db)
    except OrganizationForbiddenError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc


def _project_org_id(row: KeywordResearchProject | KeywordClusterProject) -> int:
    try:
        state = sqla_inspect(row)
        if "organization_id" in set(getattr(state, "unloaded", ()) or ()):
            oid = None
        else:
            oid = getattr(row, "organization_id", None)
    except Exception:
        oid = getattr(row, "organization_id", None)
    if oid is not None and int(oid) > 0:
        return int(oid)
    uid = getattr(row, "user_id", None)
    return personal_organization_id(int(uid)) if uid else 0


def _assert_org_project_access(
    db: Session,
    user: User,
    row: KeywordResearchProject | KeywordClusterProject | None,
    *,
    require_result: bool = False,
) -> KeywordResearchProject | KeywordClusterProject:
    if row is None:
        raise HTTPException(status_code=404, detail="Project not found")
    org_id = _project_org_id(row)
    if not user_belongs_to_organization(db, user_id=int(user.id), organization_id=org_id):
        raise HTTPException(status_code=404, detail="Project not found")
    if require_result and not getattr(row, "result_json", None):
        raise HTTPException(status_code=404, detail="Project not found")
    return row


def _assert_workspace_write(db: Session, user: User, org_id: int) -> None:
    assert_permission(user, PERM_WORKSPACE_WRITE, db=db, organization_id=org_id)


_PHYS_COL_CACHE: dict[tuple[int, str], set[str]] = {}


def _physical_columns(db: Session, table: str) -> set[str]:
    bind = db.get_bind()
    key = (id(bind), table)
    hit = _PHYS_COL_CACHE.get(key)
    if hit is not None:
        return hit
    try:
        insp = sqla_inspect(bind)
        if table not in set(insp.get_table_names() or []):
            cols: set[str] = set()
        else:
            cols = {str(c["name"]) for c in insp.get_columns(table)}
    except Exception:
        cols = set()
    _PHYS_COL_CACHE[key] = cols
    return cols


def _orm_load_only_physical(db: Session, model):
    phys = _physical_columns(db, str(model.__table__.name))
    mapper = class_mapper(model)
    if phys:
        cols = [getattr(model, attr.key) for attr in mapper.column_attrs if attr.key in phys]
    else:
        cols = [getattr(model, attr.key) for attr in mapper.column_attrs if attr.key != "organization_id"]
    return load_only(*cols) if cols else None


def _insert_kw_project(db: Session, model, fields: dict[str, Any]) -> int:
    """INSERT only columns that exist in the live DB (PG may lack organization_id)."""
    table = model.__table__
    phys = _physical_columns(db, str(table.name))
    if phys:
        values = {k: v for k, v in fields.items() if k in phys and k in table.c}
    else:
        values = {k: v for k, v in fields.items() if k in table.c and k != "organization_id"}
    result = db.execute(sqla_insert(table).values(**values).returning(table.c.id))
    pid = result.scalar()
    if pid is None:
        raise RuntimeError(f"insert {table.name} returned no id")
    return int(pid)


def _get_kw_project(db: Session, model, project_id: int):
    q = db.query(model).filter(model.id == int(project_id))
    opts = _orm_load_only_physical(db, model)
    if opts is not None:
        q = q.options(opts)
    return q.first()


def _list_kw_projects(db: Session, model, *, user: User, org_id: int, limit: int):
    phys = _physical_columns(db, str(model.__table__.name))
    q = db.query(model)
    opts = _orm_load_only_physical(db, model)
    if opts is not None:
        q = q.options(opts)
    if phys and "organization_id" in phys:
        q = q.filter(model.organization_id == org_id)
    else:
        q = q.filter(model.user_id == user.id)
    return q.order_by(desc(model.created_at)).limit(limit).all()


def _country_label(code: str) -> str:
    c = (code or "").lower()
    return {
        "vn": "Vietnam",
        "us": "United States",
        "gb": "United Kingdom",
    }.get(c, (code or "").upper() if len(c) == 2 else (code or ""))


class ResearchRunBody(BaseModel):
    seeds: list[str] = Field(..., min_length=1, max_length=50)
    domain: str | None = Field(default=None, max_length=200)
    url: str | None = Field(default=None, max_length=2000)
    engine: str = Field(default="google", max_length=32)
    language: str = Field(default="vi", max_length=16)
    country: str = Field(default="vn", max_length=16)
    device: str = Field(default="desktop", max_length=16)
    cluster_mode: str | None = Field(default=None, max_length=16, description="hybrid | tfidf | both")
    cluster_strictness: str | None = Field(default=None, max_length=16, description="strict | normal | loose")
    cluster_fetch_serp: bool | None = Field(default=None, description="Hybrid SERP fetch; None = env default")
    cluster_max_keywords: int | None = Field(default=None, description="Giới hạn KW cho clustering; server sẽ clamp an toàn")
    max_keywords: int | None = Field(default=None, ge=20, le=2000, description="Giới hạn KW trả về sau expansion")
    deep_analysis: bool = Field(default=False, description="Bật SERP/KD/refresh volume (tốn thêm credit)")


class ClusterSaveBody(BaseModel):
    keywords: list[str] = Field(..., min_length=1, max_length=5000)
    project_name: str | None = Field(default=None, max_length=255)
    fetch_serp: bool = Field(default=True)
    country: str = Field(default="vn", max_length=8)
    language: str = Field(default="vi", max_length=8)
    device: str = Field(default="desktop", max_length=16)
    cluster_strictness: str = Field(default="normal", max_length=16)
    url: str | None = Field(default=None, max_length=2000)
    result: dict = Field(..., description="Raw clustering response payload")


class VolumeBatchBody(BaseModel):
    keywords: list[str] = Field(..., min_length=1, max_length=5000)
    country: str = Field(default="vn", max_length=8)
    language: str = Field(default="vi", max_length=8)


class SuggestBody(BaseModel):
    seeds: list[str] = Field(..., min_length=1, max_length=10)
    country: str = Field(default="vn", max_length=8)
    language: str = Field(default="vi", max_length=8)
    max_keywords: int = Field(default=120, ge=10, le=300)


@router.post("/volume/batch")
def api_keyword_volume_batch(
    body: VolumeBatchBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active_trial),
) -> JSONResponse:
    from app.services.dataforseo_billing_service import (
        DataforseoStats,
        maybe_log_dataforseo_usage,
        preview_volume_batch_billing,
        user_may_use_dataforseo_live,
    )
    from app.services.cau_ai_fix_entitlement import resolve_cau_ai_plan_slug

    kws = [str(s).strip() for s in body.keywords if str(s).strip()]
    if not kws:
        raise HTTPException(status_code=400, detail="At least one keyword required")

    billing_user = current_user
    feature = feature_key_for("keyword_research")
    enforce_saas_access(db, billing_user, feature_code=feature, quantity=max(1, (len(kws) + 99) // 100))

    stats: dict[str, Any] = {}
    allow_live = user_may_use_dataforseo_live(db, billing_user)
    preview = preview_volume_batch_billing(
        db,
        billing_user,
        keywords=kws,
        country=body.country,
        language=body.language,
    )
    plan_slug = resolve_cau_ai_plan_slug(db, billing_user)
    credit_cost = int(preview.credit_cost or 0)
    job_id = build_job_id("volume_batch", billing_user.id, uuid.uuid4().hex[:10])
    gate_req = CostGateRequest(
        db=db,
        user=billing_user,
        provider=ExternalApiProvider.DATAFORSEO,
        action="volume_batch",
        credit_cost=0,
        feature_key=feature,
        quota_units=max(1, (len(kws) + 99) // 100),
        require_paid_plan=True,
        skip_credit_check=True,
        single_flight_key=f"volume_batch:{billing_user.id}",
    )
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=max(1, (len(kws) + 99) // 100),
        job_id=job_id,
        idempotency_key=job_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            with cost_gate_session(gate_req):
                rows = enrich_keyword_volumes_cached(
                    kws,
                    country=body.country,
                    language=body.language,
                    allow_api=allow_live,
                    stats=stats,
                )
        except Exception:
            sgate.release(reason="job_failed")
            raise

        out: dict[str, dict] = {}
        for r in rows or []:
            kw = str(r.get("keyword") or "").strip()
            if not kw:
                continue
            out[kw] = {
                "search_volume": int(r.get("search_volume") or 0),
                "volume_source": str(r.get("volume_source") or ""),
                "confidence": float(r.get("confidence") or 0.0),
            }

        sgate.commit()
        dfs_stats = DataforseoStats()
        dfs_stats.merge(stats)
        maybe_log_dataforseo_usage(
            db,
            stats=dfs_stats,
            user_id=billing_user.id,
            action="volume_batch",
            endpoint="search_volume/live",
            api_mode="live",
            credit_charged=credit_cost,
            cache_hit=int(preview.cache_hits),
            keyword_count=len(kws),
            plan_slug=plan_slug,
        )
        db.commit()

    return JSONResponse(content={"items": out, "billing": preview.to_dict() if preview else None})


@router.post("/suggest")
def api_keyword_suggest(
    body: SuggestBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_active_trial),
) -> JSONResponse:
    """Keyword Suggestion — Google Ads primary, DataForSEO fallback (auth + credit + tenant)."""
    from app.services.dataforseo_billing_service import (
        DataforseoStats,
        maybe_log_dataforseo_usage,
        preview_research_billing,
        user_may_use_dataforseo_live,
    )
    from app.services.keyword_ads_expand import expand_keyword_ideas

    seeds = [str(s).strip() for s in body.seeds if str(s).strip()][:10]
    if not seeds:
        raise HTTPException(status_code=400, detail="At least one seed keyword required")

    org_id = _kw_org_id(db, current_user)
    _assert_workspace_write(db, current_user, org_id)

    feature = feature_key_for("keyword_research")
    enforce_saas_access(
        db,
        current_user,
        feature_code=feature,
        quantity=max(1, len(seeds)),
        organization_id=org_id,
    )
    allow_live = user_may_use_dataforseo_live(db, current_user)
    preview = preview_research_billing(
        db,
        current_user,
        seeds=seeds,
        country=body.country,
        language=body.language,
    )
    from_cache = int(preview.cache_hits or 0) >= len(seeds) and int(preview.api_tasks or 0) == 0
    credit_cost = cost_keyword_suggest(db=db, from_cache=from_cache) if credits_enforced() else 0
    job_id = build_job_id("keyword_suggest", current_user.id, org_id, uuid.uuid4().hex[:10])
    gate_req = CostGateRequest(
        db=db,
        user=current_user,
        provider=ExternalApiProvider.DATAFORSEO,
        action="keyword_suggest",
        credit_cost=0,
        feature_key=feature,
        quota_units=max(1, len(seeds)),
        require_paid_plan=True,
        skip_credit_check=True,
        single_flight_key=f"keyword_suggest:{current_user.id}",
    )
    with saas_action_gate(
        db,
        current_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=max(1, len(seeds)),
        job_id=job_id,
        idempotency_key=job_id,
        organization_id=org_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            with cost_gate_session(gate_req):
                stats: dict[str, Any] = {}
                rows = expand_keyword_ideas(
                    seeds,
                    country=body.country,
                    language=body.language,
                    total_cap=body.max_keywords,
                    allow_api=allow_live,
                    stats=stats,
                )
        except Exception:
            sgate.release(reason="job_failed")
            raise

        keywords = [
            {
                "keyword": str(r.get("keyword") or "").strip(),
                "search_volume": int(r.get("search_volume") or 0),
                "source": str(r.get("source") or "dataforseo"),
                "cpc_avg": r.get("cpc_avg"),
                "competition_index": r.get("competition_index"),
            }
            for r in (rows or [])
            if str(r.get("keyword") or "").strip()
        ]
        sgate.commit()
        sgate.record_usage(metadata={"seeds": len(seeds), "keywords": len(keywords), "organization_id": org_id})
        dfs_stats = DataforseoStats()
        dfs_stats.merge(stats)
        from app.services.cau_ai_fix_entitlement import resolve_cau_ai_plan_slug

        maybe_log_dataforseo_usage(
            db,
            stats=dfs_stats,
            user_id=current_user.id,
            action="keyword_suggest",
            endpoint="keywords_for_keywords/live",
            api_mode=preview.api_mode,
            credit_charged=credit_cost,
            cache_hit=dfs_stats.cache_hits,
            keyword_count=len(seeds),
            plan_slug=resolve_cau_ai_plan_slug(db, current_user),
        )
        db.commit()

    return JSONResponse(
        content={
            "keywords": keywords,
            "meta": {
                "seeds": seeds,
                "count": len(keywords),
                "organization_id": org_id,
                "billing": preview.to_dict() if preview else None,
            },
        }
    )


def _cluster_label(keywords: list[str]) -> str:
    xs = [str(s).strip() for s in (keywords or []) if str(s).strip()]
    if not xs:
        return "Keyword clustering"
    head = " + ".join(xs[:8])
    if len(xs) > 8:
        head += " + …"
    return head[:2000]


def _normalize_cluster_project_name(name: str | None, keyword_count: int) -> str:
    collapsed = " ".join(str(name or "").split())
    if collapsed:
        return collapsed[:255]
    from datetime import datetime

    stamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    return f"Gom nhóm {stamp}"[:255]


@router.post("/cluster/save")
def api_cluster_save(
    body: ClusterSaveBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    kws = [str(s).strip() for s in body.keywords if str(s).strip()]
    if not kws:
        raise HTTPException(status_code=400, detail="At least one keyword required")
    strict = (body.cluster_strictness or "normal").strip().lower()
    if strict not in ("strict", "normal", "loose"):
        strict = "normal"
    payload = body.result if isinstance(body.result, dict) else {}
    pname = _normalize_cluster_project_name(body.project_name, len(kws))
    org_id = _kw_org_id(db, current_user)
    _assert_workspace_write(db, current_user, org_id)
    credit_cost = cost_cluster_save(db=db) if credits_enforced() else 0
    job_id = build_job_id("cluster_save", billing_user.id, org_id, uuid.uuid4().hex[:10])
    with saas_action_gate(
        db,
        billing_user,
        feature_code="keyword_cluster",
        credit_cost=credit_cost,
        quota_units=1,
        job_id=job_id,
        idempotency_key=job_id,
        organization_id=org_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            pid = _insert_kw_project(
                db,
                KeywordClusterProject,
                {
                    "user_id": current_user.id,
                    "organization_id": org_id,
                    "project_name": pname,
                    "keywords_label": pname,
                    "keywords_json": json.dumps(kws, ensure_ascii=False),
                    "language": (body.language or "vi").strip()[:16],
                    "country": (body.country or "vn").strip()[:16],
                    "device": (body.device or "desktop").strip()[:16],
                    "cluster_strictness": strict[:16],
                    "fetch_serp": 1 if bool(body.fetch_serp) else 0,
                    "brand_url": (body.url or "").strip()[:2000],
                    "result_json": json.dumps(payload, ensure_ascii=False),
                },
            )
            sgate.commit()
            sgate.record_usage(metadata={"kw_count": len(kws)})
            db.commit()
            return JSONResponse(content={"ok": True, "project_id": pid, "project_name": pname})
        except Exception:
            sgate.release(reason="cluster_save_failed")
            db.rollback()
            raise


@router.get("/cluster/history")
def api_cluster_history(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    organization_id: int | None = Query(default=None),
    limit: int = Query(50, ge=1, le=200),
) -> JSONResponse:
    org_id = _kw_org_id(db, current_user, organization_id)
    rows = _list_kw_projects(
        db, KeywordClusterProject, user=current_user, org_id=org_id, limit=limit
    )
    out: list[dict] = []
    for r in rows:
        kw_count = 0
        try:
            kw_count = len(json.loads(r.keywords_json or "[]"))
        except (json.JSONDecodeError, TypeError):
            kw_count = 0
        display_name = (r.project_name or "").strip() or f"Dự án #{r.id}"
        out.append(
            {
                "id": r.id,
                "project_name": display_name,
                "keywords_label": r.keywords_label,
                "keyword_count": kw_count,
                "language": r.language,
                "country": r.country,
                "device": r.device,
                "cluster_strictness": r.cluster_strictness,
                "fetch_serp": bool(r.fetch_serp),
                "brand_url": r.brand_url,
                "created_at": r.created_at.isoformat() if r.created_at else "",
            }
        )
    return JSONResponse(content={"projects": out})


@router.get("/cluster/history/{project_id}")
def api_cluster_history_item(
    project_id: int,
    summary_only: bool = Query(default=False),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    r = _get_kw_project(db, KeywordClusterProject, project_id)
    r = _assert_org_project_access(db, current_user, r, require_result=True)
    payload = {
        "project": {
            "id": r.id,
            "project_name": (r.project_name or "").strip() or f"Dự án #{r.id}",
            "keywords_label": r.keywords_label,
            "language": r.language,
            "country": r.country,
            "device": r.device,
            "cluster_strictness": r.cluster_strictness,
            "fetch_serp": bool(r.fetch_serp),
            "brand_url": r.brand_url,
            "created_at": r.created_at.isoformat() if r.created_at else "",
        },
        "result": json.loads(r.result_json),
        "keywords": json.loads(r.keywords_json) if r.keywords_json else [],
    }
    return JSONResponse(content=shape_keyword_cluster_history_response(payload, summary_only=summary_only))


@router.delete("/cluster/history/{project_id}")
def api_cluster_history_delete(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    r = _get_kw_project(db, KeywordClusterProject, project_id)
    _assert_org_project_access(db, current_user, r)
    db.delete(r)
    db.commit()
    return JSONResponse(content={"ok": True})


@router.get("/research")
def api_keyword_research(
    seed_keyword: str = Query(..., min_length=1, max_length=2000),
    domain: str | None = Query(default=None, max_length=200),
    url: str | None = Query(default=None, max_length=2000),
    engine: str = Query("google", max_length=32),
    language: str = Query("vi", max_length=16),
    country: str = Query("vn", max_length=8),
    device: str = Query("desktop", max_length=16),
    cluster_mode: str | None = Query(default=None, max_length=16),
    cluster_strictness: str | None = Query(default=None, max_length=16),
    cluster_fetch_serp: str | None = Query(default=None, max_length=8),
    cluster_max_keywords: int | None = Query(default=None),
    db: Session = Depends(get_db),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    from app.services.cau_ai_fix_entitlement import resolve_cau_ai_plan_slug
    from app.services.dataforseo_billing_service import (
        DataforseoStats,
        maybe_log_dataforseo_usage,
        preview_research_billing,
        user_may_use_dataforseo_live,
    )

    cfs: bool | None = None
    if cluster_fetch_serp is not None and str(cluster_fetch_serp).strip() != "":
        cfs = str(cluster_fetch_serp).lower() in ("1", "true", "yes", "on")

    seeds = [seed_keyword.strip()]
    feature = feature_key_for("keyword_research")
    enforce_saas_access(db, billing_user, feature_code=feature, quantity=1)
    allow_live = user_may_use_dataforseo_live(db, billing_user)
    preview = preview_research_billing(db, billing_user, seeds=seeds, country=country, language=language)
    credit_cost = int(preview.credit_cost or 0)
    job_id = build_job_id("keyword_research_get", billing_user.id, uuid.uuid4().hex[:10])
    gate_req = CostGateRequest(
        db=db,
        user=billing_user,
        provider=ExternalApiProvider.DATAFORSEO,
        action="keyword_research_get",
        credit_cost=0,
        feature_key=feature,
        quota_units=1,
        require_paid_plan=True,
        skip_credit_check=True,
        single_flight_key=f"keyword_research_get:{billing_user.id}",
    )
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=1,
        job_id=job_id,
        idempotency_key=job_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            with cost_gate_session(gate_req):
                stats: dict[str, Any] = {}
                release_db_before_long_io(db)
                payload = build_keyword_research_api_response(
                    seed_keyword=seed_keyword,
                    domain=domain,
                    url=url,
                    engine=engine,
                    language=language,
                    country=country,
                    device=device,
                    cluster_mode=cluster_mode,
                    cluster_strictness=cluster_strictness,
                    cluster_fetch_serp=cfs,
                    cluster_max_keywords=cluster_max_keywords,
                    allow_dataforseo_live=allow_live,
                    billing_stats=stats,
                )
        except CircuitOpenError as exc:
            sgate.release(reason="provider_circuit")
            raise HTTPException(
                status_code=429,
                detail={
                    "code": "keyword_provider_busy",
                    "message": (
                        "Dịch vụ nghiên cứu từ khóa đang quá tải / hết hạn mức tạm thời. "
                        f"Thử lại sau khoảng {int(getattr(exc, 'retry_after_sec', 30) or 30)} giây."
                    ),
                    "retry_after_sec": float(getattr(exc, "retry_after_sec", 30) or 30),
                },
            ) from exc
        except Exception:
            sgate.release(reason="job_failed")
            raise

        if payload.get("meta", {}).get("error"):
            sgate.release(reason="bad_request")
            return JSONResponse(status_code=400, content=payload)

        try:
            ensure_db_alive(db)
        except Exception as exc:
            sgate.release(reason="db_reconnect_failed")
            raise HTTPException(
                status_code=503,
                detail={
                    "code": "db_temporarily_unavailable",
                    "message": "Kết nối dữ liệu tạm thời gián đoạn sau khi gọi API ngoài. Vui lòng thử lại.",
                },
            ) from exc

        sgate.commit()
        dfs_stats = DataforseoStats()
        dfs_stats.merge(stats)
        maybe_log_dataforseo_usage(
            db,
            stats=dfs_stats,
            user_id=billing_user.id,
            action="keyword_research_get",
            endpoint="keywords_for_keywords/live",
            api_mode=preview.api_mode,
            credit_charged=credit_cost,
            cache_hit=dfs_stats.cache_hits,
            keyword_count=1,
            plan_slug=resolve_cau_ai_plan_slug(db, billing_user),
        )
        db.commit()
    return JSONResponse(content=payload)


@router.post("/research/run")
def api_keyword_research_run(
    body: ResearchRunBody,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    from app.services.cau_ai_fix_entitlement import resolve_cau_ai_plan_slug
    from app.services.dataforseo_billing_service import (
        DataforseoStats,
        maybe_log_dataforseo_usage,
        preview_research_billing,
        user_may_use_dataforseo_live,
    )

    seeds = [str(s).strip() for s in body.seeds if str(s).strip()][:50]
    if not seeds:
        raise HTTPException(status_code=400, detail="At least one keyword required")
    charge_user = billing_user  # require_active_trial always supplies user
    feature = feature_key_for("keyword_research")
    enforce_saas_access(db, charge_user, feature_code=feature, quantity=max(1, len(seeds)))
    allow_live = user_may_use_dataforseo_live(db, charge_user)
    preview = preview_research_billing(
        db,
        charge_user,
        seeds=seeds,
        country=body.country,
        language=body.language,
    )
    credit_cost = int(preview.credit_cost or 0)
    job_id = build_job_id("keyword_research", charge_user.id, uuid.uuid4().hex[:10])
    gate_req = CostGateRequest(
        db=db,
        user=charge_user,
        provider=ExternalApiProvider.DATAFORSEO,
        action="keyword_research",
        credit_cost=0,
        feature_key=feature,
        quota_units=max(1, len(seeds)),
        require_paid_plan=True,
        skip_credit_check=True,
        single_flight_key=f"keyword_research:{current_user.id}",
    )
    stats: dict[str, Any] = {}
    with saas_action_gate(
        db,
        charge_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=max(1, len(seeds)),
        job_id=job_id,
        idempotency_key=job_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            with cost_gate_session(gate_req):
                # Release PgBouncer server conn before long Google Ads / DFS I/O.
                release_db_before_long_io(db)
                payload = build_keyword_research_api_response(
                    seed_keywords=seeds,
                    domain=body.domain,
                    url=body.url,
                    engine=body.engine,
                    language=body.language,
                    country=body.country,
                    device=body.device,
                    cluster_mode=body.cluster_mode,
                    cluster_strictness=body.cluster_strictness,
                    cluster_fetch_serp=body.cluster_fetch_serp,
                    cluster_max_keywords=body.cluster_max_keywords,
                    max_keywords=body.max_keywords,
                    deep_analysis=body.deep_analysis,
                    allow_dataforseo_live=allow_live,
                    billing_stats=stats,
                )
        except KeywordResearchCostExceeded as exc:
            sgate.release(reason="cost_limit")
            raise HTTPException(
                status_code=429,
                detail={
                    "code": "keyword_research_cost_limit",
                    "message": str(exc),
                    "cost": (exc.stats or stats),
                },
            ) from exc
        except CircuitOpenError as exc:
            sgate.release(reason="provider_circuit")
            raise HTTPException(
                status_code=429,
                detail={
                    "code": "keyword_provider_busy",
                    "message": (
                        "Dịch vụ nghiên cứu từ khóa đang quá tải / hết hạn mức tạm thời. "
                        f"Thử lại sau khoảng {int(getattr(exc, 'retry_after_sec', 30) or 30)} giây."
                    ),
                    "retry_after_sec": float(getattr(exc, "retry_after_sec", 30) or 30),
                },
            ) from exc
        except Exception:
            sgate.release(reason="job_failed")
            raise

        if payload.get("meta", {}).get("error"):
            sgate.release(reason="bad_request")
            return JSONResponse(status_code=400, content=payload)

        try:
            ensure_db_alive(db)
        except Exception as exc:
            sgate.release(reason="db_reconnect_failed")
            raise HTTPException(
                status_code=503,
                detail={
                    "code": "db_temporarily_unavailable",
                    "message": "Kết nối dữ liệu tạm thời gián đoạn sau khi gọi API ngoài. Vui lòng thử lại.",
                },
            ) from exc

        def _persist_success() -> None:
            sgate.commit()
            dfs_stats = DataforseoStats()
            dfs_stats.merge(stats)
            record_quota_usage(
                db,
                charge_user.id,
                feature,
                quantity=max(1, len(seeds)),
                credits_used=credit_cost,
            )
            maybe_log_dataforseo_usage(
                db,
                stats=dfs_stats,
                user_id=charge_user.id,
                action="keyword_research",
                endpoint="keywords_for_keywords/live",
                api_mode=preview.api_mode,
                credit_charged=credit_cost,
                cache_hit=dfs_stats.cache_hits,
                keyword_count=len(seeds),
                plan_slug=resolve_cau_ai_plan_slug(db, charge_user),
            )
            label = " + ".join(seeds[:10])
            if len(seeds) > 10:
                label += " + …"
            project_id = _insert_kw_project(
                db,
                KeywordResearchProject,
                {
                    "user_id": current_user.id,
                    "organization_id": _kw_org_id(db, current_user),
                    "keywords_label": label[:2000],
                    "seeds_json": json.dumps(seeds, ensure_ascii=False),
                    "language": body.language,
                    "country": body.country,
                    "engine": body.engine,
                    "location_label": _country_label(body.country),
                    "status": "completed",
                    "result_json": json.dumps(payload, ensure_ascii=False),
                },
            )
            db.commit()
            payload["meta"]["project_id"] = project_id

        try:
            _persist_success()
        except OperationalError:
            # One reconnect retry after long provider call.
            try:
                ensure_db_alive(db)
                _persist_success()
            except Exception as exc:
                sgate.release(reason="persist_failed")
                raise HTTPException(
                    status_code=503,
                    detail={
                        "code": "db_temporarily_unavailable",
                        "message": "Không lưu được kết quả research vì lỗi kết nối DB. Vui lòng thử lại.",
                    },
                ) from exc
    return JSONResponse(content=payload)


@router.get("/research/history")
def api_research_history(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    organization_id: int | None = Query(default=None),
    limit: int = Query(50, ge=1, le=200),
) -> JSONResponse:
    org_id = _kw_org_id(db, current_user, organization_id)
    rows = _list_kw_projects(
        db, KeywordResearchProject, user=current_user, org_id=org_id, limit=limit
    )
    out: list[dict] = []
    for r in rows:
        out.append(
            {
                "id": r.id,
                "keywords_label": r.keywords_label,
                "language": r.language,
                "location_label": r.location_label or _country_label(r.country),
                "status": r.status,
                "created_at": r.created_at.isoformat() if r.created_at else "",
            }
        )
    return JSONResponse(content={"projects": out})


@router.get("/research/history/{project_id}")
def api_research_history_item(
    project_id: int,
    summary_only: bool = Query(default=False),
    keywords_page: int | None = Query(default=None, ge=1),
    keywords_page_size: int = Query(default=100, ge=1, le=500),
    clusters_page: int | None = Query(default=None, ge=1),
    clusters_page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    r = _get_kw_project(db, KeywordResearchProject, project_id)
    r = _assert_org_project_access(db, current_user, r, require_result=True)
    payload = json.loads(r.result_json)
    shaped = shape_keyword_research_response(
        payload,
        keywords_page=keywords_page,
        keywords_page_size=keywords_page_size,
        clusters_page=clusters_page,
        clusters_page_size=clusters_page_size,
        summary_only=summary_only,
    )
    return JSONResponse(content=shaped)


@router.delete("/research/history/{project_id}")
def api_research_history_delete(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    r = _get_kw_project(db, KeywordResearchProject, project_id)
    _assert_org_project_access(db, current_user, r)
    db.delete(r)
    db.commit()
    return JSONResponse(content={"ok": True})


@router.post("/research/history/{project_id}/delete")
def api_research_history_delete_post(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    """
    Backward-compatible delete endpoint for clients/proxies that block DELETE.
    """
    r = _get_kw_project(db, KeywordResearchProject, project_id)
    _assert_org_project_access(db, current_user, r)
    db.delete(r)
    db.commit()
    return JSONResponse(content={"ok": True})


@router.post("/research/delete/{project_id}")
def api_research_delete_alias(
    project_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    """
    Extra alias for compatibility with legacy clients.
    """
    r = _get_kw_project(db, KeywordResearchProject, project_id)
    _assert_org_project_access(db, current_user, r)
    db.delete(r)
    db.commit()
    return JSONResponse(content={"ok": True})


@router.post("/research/run/async")
def api_keyword_research_run_async(
    body: ResearchRunBody,
    db: Session = Depends(get_db),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    """Queue full research + clustering on Celery (``KEYWORD_RESEARCH_ASYNC_CELERY=1`` + worker)."""
    if os.getenv("KEYWORD_RESEARCH_ASYNC_CELERY", "0").lower() not in ("1", "true", "yes"):
        raise HTTPException(
            status_code=503,
            detail="Async keyword research disabled. Set KEYWORD_RESEARCH_ASYNC_CELERY=1 and run a Celery worker.",
        )
    try:
        from app.workers.keyword_tasks import run_keyword_research_task
    except ImportError as exc:
        raise HTTPException(status_code=503, detail=f"Celery worker module unavailable: {exc}") from exc
    seeds = [str(s).strip() for s in body.seeds if str(s).strip()][:50]
    if not seeds:
        raise HTTPException(status_code=400, detail="At least one keyword required")
    from app.services.dataforseo_billing_service import preview_research_billing

    feature = feature_key_for("keyword_research")
    enforce_saas_access(db, billing_user, feature_code=feature, quantity=max(1, len(seeds)))
    preview = preview_research_billing(
        db,
        billing_user,
        seeds=seeds,
        country=body.country,
        language=body.language,
    )
    credit_cost = int(preview.credit_cost or 0)
    job_id = build_job_id("keyword_research_async", billing_user.id, uuid.uuid4().hex[:12])
    gate_req = CostGateRequest(
        db=db,
        user=billing_user,
        provider=ExternalApiProvider.DATAFORSEO,
        action="keyword_research_async",
        credit_cost=0,
        feature_key=feature,
        quota_units=max(1, len(seeds)),
        require_paid_plan=True,
        skip_credit_check=True,
        single_flight_key=f"keyword_research_async:{billing_user.id}",
    )
    with cost_gate_session(gate_req):
        pass
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=max(1, len(seeds)),
        job_id=job_id,
        idempotency_key=job_id,
        organization_id=_kw_org_id(db, billing_user),
        charge_mode="reserve",
    ) as sgate:
        try:
            owner_id = int(billing_user.id)
            payload = body.model_dump()
            payload["_owner_user_id"] = owner_id
            payload["_credit_reservation_id"] = sgate.reservation_id
            payload["_credit_job_id"] = job_id
            task = run_keyword_research_task.delay(payload)
            from app.services.tenant_org import personal_organization_id

            register_task_owner(
                task.id,
                owner_id,
                organization_id=personal_organization_id(owner_id),
                db=db,
            )
            # Keep reservation status=reserved until worker success; persist hold only.
            db.commit()
        except Exception:
            sgate.release(reason="enqueue_failed")
            db.commit()
            raise
        return JSONResponse(
            content={
                "task_id": task.id,
                "state": "QUEUED",
                "poll_url": f"/keywords/research/task/{task.id}",
            }
        )


@router.get("/research/task/{task_id}")
def api_keyword_research_task_status(
    task_id: str,
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    """Poll Celery task result for ``/research/run/async``."""
    role = normalize_role(current_user.role)
    assert_task_owner(task_id, current_user.id, role=role)
    try:
        from celery.result import AsyncResult

        from app.queue.celery_app import celery_app
    except ImportError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    r = AsyncResult(task_id, app=celery_app)
    if r.state == "PENDING":
        return JSONResponse(content={"task_id": task_id, "state": r.state})
    if r.state == "SUCCESS":
        return JSONResponse(content={"task_id": task_id, "state": r.state, "result": r.result})
    return JSONResponse(content={"task_id": task_id, "state": r.state, "info": str(r.info)})


@router.post("/cluster")
async def api_keyword_cluster(
    request: Request,
    db: Session = Depends(get_db),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    """SERP-primary clustering: merge keyword pairs with >=3 shared top-10 Google URLs."""
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON object required")
    kws = body.get("keywords") or body.get("keyword_list")
    if not isinstance(kws, list) or not kws:
        raise HTTPException(status_code=400, detail="keywords array required")
    strings = [str(x).strip() for x in kws if str(x).strip()]
    fetch_serp = body.get("fetch_serp")
    fetch_on = True if fetch_serp is None else bool(fetch_serp)
    brand_hint = str(body.get("url") or body.get("domain") or "").strip() or None
    serp_country = str(body.get("country") or "").strip() or None
    serp_language = str(body.get("language") or "").strip() or None
    serp_device = str(body.get("device") or "").strip() or None
    strict = body.get("cluster_strictness")
    strict_s = str(strict).strip().lower() if strict is not None else None
    feature = "keyword_cluster"
    credit_cost = cost_cluster_sync(db=db, keyword_count=len(strings), fetch_serp=fetch_on)
    job_id = build_job_id("keyword_cluster", billing_user.id, uuid.uuid4().hex[:10])
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=1,
        job_id=job_id,
        idempotency_key=job_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            payload = build_keyword_cluster_api_response(
                strings,
                fetch_serp=fetch_on,
                brand_host_hint=brand_hint,
                serp_country=serp_country,
                serp_language=serp_language,
                serp_device=serp_device,
                cluster_strictness=strict_s,
            )
        except Exception:
            sgate.release(reason="job_failed")
            raise
        sgate.commit()
        db.commit()
    return JSONResponse(content=payload)


@router.post("/cluster/async")
async def api_keyword_cluster_async(
    request: Request,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    """
    Run clustering in background and expose progress % via polling.

    Returns: { job_id, state, progress, poll_url }
    """
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON object required")
    kws = body.get("keywords") or body.get("keyword_list")
    if not isinstance(kws, list) or not kws:
        raise HTTPException(status_code=400, detail="keywords array required")
    strings = [str(x).strip() for x in kws if str(x).strip()]
    fetch_serp = body.get("fetch_serp")
    fetch_on = True if fetch_serp is None else bool(fetch_serp)
    brand_hint = str(body.get("url") or body.get("domain") or "").strip() or None
    serp_country = str(body.get("country") or "").strip() or None
    serp_language = str(body.get("language") or "").strip() or None
    serp_device = str(body.get("device") or "").strip() or None
    strict = body.get("cluster_strictness")
    strict_s = str(strict).strip().lower() if strict is not None else None

    from app.services.job_store import cleanup_expired, create_job, ensure_job_schema, get_job, mark_stale_jobs_failed, mark_stale_queued_failed, update_job

    ensure_job_schema()
    cleanup_expired(ttl_seconds=int(os.getenv("JOB_TTL_SECONDS", "1800")))
    mark_stale_jobs_failed(stale_seconds=int(os.getenv("JOB_WATCHDOG_SECONDS", "300")))
    mark_stale_queued_failed(stale_seconds=int(os.getenv("JOB_QUEUE_STALE_SECONDS", "90")))

    org_id = _kw_org_id(db, billing_user)
    job_payload = {
        "user_id": current_user.id,
        "organization_id": org_id,
        "keywords": strings,
        "fetch_serp": fetch_on,
        "brand_host_hint": brand_hint,
        "serp_country": serp_country,
        "serp_language": serp_language,
        "serp_device": serp_device,
        "cluster_strictness": strict_s,
    }
    job = create_job(
        job_type="keyword_cluster",
        message="Queued",
        payload=job_payload,
        organization_id=org_id,
        user_id=int(current_user.id),
    )
    feature = "keyword_cluster"
    credit_cost = cost_cluster_async(db=db, keyword_count=len(strings), fetch_serp=fetch_on)
    job_credit_id = build_job_id("keyword_cluster_async", billing_user.id, job.job_id)
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=1,
        job_id=job_credit_id,
        idempotency_key=job_credit_id,
        organization_id=org_id,
        charge_mode="reserve",
    ) as sgate:
        job_payload["credit_job_id"] = job_credit_id
        job_payload["billing_user_id"] = int(billing_user.id)
        job_payload["reservation_id"] = sgate.reservation_id
        db.commit()

    inline = os.getenv("KEYWORD_CLUSTER_INLINE", "1").lower() in ("1", "true", "yes")
    if inline:
        from app.services.cluster_job_runner import run_cluster_job

        threading.Thread(
            target=run_cluster_job,
            args=(job.job_id, job_payload),
            daemon=True,
            name=f"cluster-{job.job_id[:8]}",
        ).start()

    st0 = get_job(job.job_id)
    return JSONResponse(
        content={
            "job_id": job.job_id,
            "state": st0.state if st0 else job.state,
            "progress": st0.progress if st0 else job.progress,
            "poll_url": f"/keywords/cluster/job/{job.job_id}",
        }
    )


@router.get("/cluster/job/{job_id}")
def api_keyword_cluster_job(
    job_id: str,
    current_user: User = Depends(get_current_user),
) -> JSONResponse:
    from app.services.job_store import cleanup_expired, ensure_job_schema, get_job, mark_stale_jobs_failed, mark_stale_queued_failed

    ensure_job_schema()
    cleanup_expired(ttl_seconds=int(os.getenv("JOB_TTL_SECONDS", "1800")))
    mark_stale_jobs_failed(stale_seconds=int(os.getenv("JOB_WATCHDOG_SECONDS", "300")))
    mark_stale_queued_failed(stale_seconds=int(os.getenv("JOB_QUEUE_STALE_SECONDS", "90")))
    st = get_job_for_user(job_id, user_id=current_user.id)
    if not st:
        raise HTTPException(status_code=404, detail="Không tìm thấy job.")
    out = {
        "job_id": st.job_id,
        "state": st.state,
        "progress": st.progress,
        "message": st.message,
    }
    if st.state == "SUCCESS" and st.result is not None:
        out["result"] = st.result
    if st.state == "ERROR":
        out["error"] = st.error or "Unknown error"
    return JSONResponse(content=out)


@router.post("/import/excel")
async def api_import_keywords_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    billing_user: User = Depends(require_active_trial),
) -> JSONResponse:
    """
    Upload an Excel (.xlsx) and extract keyword list.

    Heuristics:
    - Detect a 'keyword' column from header (keyword/từ khóa/query/kw)
    - Else use first column
    """
    name = (file.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Only .xlsx supported")
    raw = await file.read()
    if not raw or len(raw) < 50:
        raise HTTPException(status_code=400, detail="Empty file")
    if len(raw) > 12_000_000:
        raise HTTPException(status_code=413, detail="File too large (max 12MB)")
    try:
        import openpyxl

        wb = openpyxl.load_workbook(filename=BytesIO(raw), data_only=True, read_only=True)
        ws = wb.active
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid xlsx: {exc}") from exc

    credit_cost = cost_import_excel(db=db) if credits_enforced() else 0
    feature = feature_key_for("keyword_research")
    job_id = build_job_id("import_excel", billing_user.id, uuid.uuid4().hex[:10])
    with saas_action_gate(
        db,
        billing_user,
        feature_code=feature,
        credit_cost=credit_cost,
        quota_units=1,
        job_id=job_id,
        idempotency_key=job_id,
        charge_mode="reserve",
    ) as sgate:
        try:
            def norm(x: str) -> str:
                return (x or "").strip().lower()

            header_row_idx: int | None = None
            kw_col_idx: int | None = None  # 1-based
            header_needles = ("keyword", "từ khóa", "tu khoa", "query", "kw")

            for r_idx in range(1, 6):
                row = [ws.cell(row=r_idx, column=c).value for c in range(1, min(30, ws.max_column or 1) + 1)]
                vals = [norm(str(v)) for v in row if v is not None]
                if not vals:
                    continue
                for c_idx in range(1, len(row) + 1):
                    v = row[c_idx - 1]
                    if v is None:
                        continue
                    nv = norm(str(v))
                    if any(n in nv for n in header_needles):
                        header_row_idx = r_idx
                        kw_col_idx = c_idx
                        break
                if kw_col_idx is not None:
                    break

            if kw_col_idx is None:
                kw_col_idx = 1
                header_row_idx = 1

            out: list[str] = []
            seen: set[str] = set()
            max_rows = 5000
            start_row = (header_row_idx or 1) + 1

            for r_idx in range(start_row, min(start_row + max_rows, (ws.max_row or start_row) + 1)):
                v = ws.cell(row=r_idx, column=kw_col_idx).value
                if v is None:
                    continue
                s = str(v).strip()
                if not s:
                    continue
                if len(s) > 250:
                    s = s[:250].strip()
                low = s.lower()
                if r_idx == start_row and any(n in low for n in header_needles):
                    continue
                if low in seen:
                    continue
                seen.add(low)
                out.append(s)
                if len(out) >= 2000:
                    break

            sgate.commit()
            sgate.record_usage(metadata={"rows": len(out)})
            return JSONResponse(
                content={
                    "keywords": out,
                    "meta": {"count": len(out), "truncated": len(out) >= 2000, "column_index": kw_col_idx},
                }
            )
        except Exception:
            sgate.release(reason="import_excel_failed")
            raise


@router.post("/export/excel")
async def api_export_excel(request: Request) -> StreamingResponse:
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON object required")
    keywords = body.get("keywords") if isinstance(body.get("keywords"), list) else []
    clusters = body.get("clusters") if isinstance(body.get("clusters"), list) else []
    if not keywords and not clusters:
        raise HTTPException(status_code=400, detail="Provide keywords and/or clusters")
    data = export_to_excel(keywords=keywords, clusters=clusters)
    name = "keyword_export.xlsx"
    return StreamingResponse(
        iter([data]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.post("/export/cluster-flat-excel")
async def api_export_cluster_flat_excel(request: Request) -> StreamingResponse:
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON object required")
    rows = body.get("rows") if isinstance(body.get("rows"), list) else []
    if not rows:
        raise HTTPException(status_code=400, detail="Provide rows")
    data = export_cluster_flat_to_excel(rows=rows)
    name = "keyword_cluster_flat.xlsx"
    return StreamingResponse(
        iter([data]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@router.post("/export/sheets")
async def api_export_sheets(request: Request) -> JSONResponse:
    body = await request.json()
    if not isinstance(body, dict):
        raise HTTPException(status_code=400, detail="JSON object required")
    keywords = body.get("keywords") if isinstance(body.get("keywords"), list) else []
    clusters = body.get("clusters") if isinstance(body.get("clusters"), list) else []
    title = str(body.get("title") or "").strip() or None
    url, err = export_to_google_sheets(keywords=keywords, clusters=clusters, title=title)
    if err:
        return JSONResponse(status_code=503, content={"ok": False, "detail": err, "spreadsheet_url": None})
    return JSONResponse(content={"ok": True, "spreadsheet_url": url})


@router.get("", response_class=RedirectResponse)
def keyword_hub_page() -> RedirectResponse:
    return RedirectResponse(url="/keywords/tools/research", status_code=302)


@router.get("/tools/research")
def keyword_research_page(
    request: Request,
    tab: str | None = Query(None),
):
    if tab == "group":
        return RedirectResponse(url="/keywords/tools/clustering", status_code=302)
    return templates.TemplateResponse(
        request=request,
        name="keyword_research.html",
        context={"nav_active": "keyword_research", "kw_tool_active": "research"},
    )


@router.get("/tools/clustering", response_class=HTMLResponse)
def keyword_clustering_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request=request,
        name="keyword_clustering.html",
        context={"nav_active": "keyword_clustering", "kw_tool_active": "clustering"},
    )


@router.get("/tools/rank-tracking", response_class=HTMLResponse)
def keyword_rank_tracking_page(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        request=request,
        name="rank_tracking.html",
        context={"nav_active": "keyword", "kw_tool_active": "rank_tracking"},
    )
